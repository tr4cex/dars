import csv
import io
import json
import re
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from xml.etree import ElementTree as ET

import httpx
import pymupdf
from bs4 import BeautifulSoup

from . import config, llm

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff", ".heic"}
TEXT_EXT = {".txt", ".md", ".markdown", ".rst", ".log", ".tex", ".srt", ".vtt", ".yaml", ".yml", ".toml", ".ini", ".xml"}
CODE_EXT = {".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".h", ".cpp", ".hpp", ".cs", ".go", ".rs", ".php", ".rb", ".sql", ".sh", ".kt", ".swift", ".css"}
HTML_EXT = {".html", ".htm"}

FORMATS = [
    {"id": "pdf", "label": "PDF", "ext": [".pdf"], "note": "Text or scanned pages"},
    {"id": "docx", "label": "Word", "ext": [".docx"], "note": "Paragraphs and tables"},
    {"id": "pptx", "label": "PowerPoint", "ext": [".pptx"], "note": "Slides and speaker notes"},
    {"id": "xlsx", "label": "Excel", "ext": [".xlsx", ".xlsm"], "note": "Every sheet"},
    {"id": "csv", "label": "CSV", "ext": [".csv", ".tsv"], "note": "Tabular data"},
    {"id": "image", "label": "Images", "ext": sorted(IMAGE_EXT), "note": "Text read by vision"},
    {"id": "epub", "label": "EPUB", "ext": [".epub"], "note": "E-books"},
    {"id": "html", "label": "HTML", "ext": sorted(HTML_EXT), "note": "Saved web pages"},
    {"id": "text", "label": "Text", "ext": sorted(TEXT_EXT | {".json"}), "note": "txt, md, json and more"},
    {"id": "code", "label": "Code", "ext": sorted(CODE_EXT), "note": "Source files"},
]

SUPPORTED = {e for f in FORMATS for e in f["ext"]}


def kind_of(path: Path):
    ext = path.suffix.lower()
    for f in FORMATS:
        if ext in f["ext"]:
            return f["id"]
    return None


def read_text_file(path: Path):
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1256", "latin-1"):
        try:
            return raw.decode(enc)
        except Exception:
            continue
    return raw.decode("utf-8", errors="ignore")


def sectionize(text, size=2600, prefix="§"):
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    paragraphs = text.split("\n\n")
    out, buf, n = [], "", 1
    for p in paragraphs:
        if buf and len(buf) + len(p) > size:
            out.append((f"{prefix}{n}", buf.strip()))
            n += 1
            buf = ""
        while len(p) > size * 1.5:
            cut = p.rfind(" ", 0, size)
            cut = cut if cut > size // 2 else size
            out.append((f"{prefix}{n}", (buf + "\n\n" + p[:cut]).strip()))
            n += 1
            buf = ""
            p = p[cut:]
        buf += ("\n\n" if buf else "") + p
    if buf.strip():
        out.append((f"{prefix}{n}", buf.strip()))
    return out


def arabic_garbled(text):
    if not text:
        return False
    forms = sum(1 for ch in text if "ﭐ" <= ch <= "﷿" or "ﹰ" <= ch <= "﻿")
    return forms / max(len(text), 1) > 0.15


def extract_pdf(path, progress):
    from . import db

    batch = db.get_setting("ocr_batch", config.DEFAULT_SETTINGS["ocr_batch"])
    scale = db.get_setting("ocr_scale", config.DEFAULT_SETTINGS["ocr_scale"])
    doc = pymupdf.open(path)
    total = len(doc)
    texts, need_ocr = {}, {}
    for i, page in enumerate(doc):
        t = page.get_text("text").strip()
        if len(t) < 30 or arabic_garbled(t):
            pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale))
            need_ocr[i] = pix.tobytes("png")
            texts[i] = t if not arabic_garbled(t) else ""
        else:
            texts[i] = t
    done = [total - len(need_ocr)]
    progress(done[0], total)
    order = sorted(need_ocr)
    groups = [order[g:g + batch] for g in range(0, len(order), batch)]

    def work(group):
        try:
            res = llm.ocr_pages([need_ocr[i] for i in group])
        except Exception:
            res = [""] * len(group)
        done[0] += len(group)
        progress(min(done[0], total), total)
        return group, res

    if groups:
        with ThreadPoolExecutor(max_workers=config.OCR_WORKERS) as ex:
            for group, res in ex.map(work, groups):
                for i, text in zip(group, res):
                    if text:
                        texts[i] = text
    doc.close()
    return [(f"p.{i + 1}", texts[i]) for i in range(total) if texts.get(i, "").strip()]


def extract_docx(path, progress):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    lines = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            m = re.search(r"heading (\d)", style)
            if m:
                lines.append("#" * int(m.group(1)) + " " + t)
            elif "title" in style:
                lines.append("# " + t)
            else:
                lines.append(t)
        elif tag == "tbl":
            tbl = Table(child, doc)
            for row in tbl.rows:
                cells = []
                for c in row.cells:
                    v = c.text.strip().replace("\n", " ")
                    if not cells or cells[-1] != v:
                        cells.append(v)
                lines.append(" | ".join(cells))
    progress(1, 1)
    return sectionize("\n\n".join(lines))


def extract_pptx(path, progress):
    from pptx import Presentation

    prs = Presentation(path)
    segs = []
    total = len(prs.slides)
    for n, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                if t.strip():
                    parts.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("Notes: " + notes)
        if parts:
            segs.append((f"Slide {n}", "\n".join(parts)))
        progress(n, total)
    return segs


def extract_xlsx(path, progress):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    segs = []
    sheets = wb.sheetnames
    for si, name in enumerate(sheets, 1):
        ws = wb[name]
        rows, header = [], None
        for r in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v).strip() for v in r]
            if not any(vals):
                continue
            if header is None:
                header = vals
            rows.append(" | ".join(vals))
        if not rows:
            continue
        for g in range(0, len(rows), 50):
            block = rows[g:g + 50]
            head = rows[0] if g else ""
            body = "\n".join(block)
            if g and head:
                body = head + "\n" + body
            segs.append((f"{name} r{g + 1}-{g + len(block)}", body))
        progress(si, len(sheets))
    wb.close()
    return segs


def extract_csv(path, progress):
    raw = read_text_file(path)
    delim = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        delim = csv.Sniffer().sniff(raw[:4000], delimiters=",;\t|").delimiter
    except Exception:
        pass
    rows = [" | ".join(c.strip() for c in r) for r in csv.reader(io.StringIO(raw), delimiter=delim) if any(c.strip() for c in r)]
    segs = []
    for g in range(0, len(rows), 60):
        block = rows[g:g + 60]
        body = "\n".join(([rows[0]] if g else []) + block)
        segs.append((f"rows {g + 1}-{g + len(block)}", body))
    progress(1, 1)
    return segs


def html_to_text(html):
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "nav", "footer", "header", "form"]):
        tag.decompose()
    title = soup.title.get_text(strip=True) if soup.title else ""
    for h in soup.find_all(re.compile(r"^h[1-6]$")):
        h.insert_before("\n\n" + "#" * int(h.name[1]) + " ")
        h.append("\n")
    for br in soup.find_all(["br", "p", "li", "tr", "div"]):
        br.append("\n")
    text = soup.get_text(" ")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return title, text.strip()


def extract_html(path, progress):
    _, text = html_to_text(read_text_file(path))
    progress(1, 1)
    return sectionize(text)


def extract_epub(path, progress):
    z = zipfile.ZipFile(path)
    container = ET.fromstring(z.read("META-INF/container.xml"))
    ns = {"c": "urn:oasis:names:tc:opendocument:xmlns:container"}
    opf_path = container.find(".//c:rootfile", ns).attrib["full-path"]
    opf = ET.fromstring(z.read(opf_path))
    base = opf_path.rsplit("/", 1)[0] + "/" if "/" in opf_path else ""
    manifest = {}
    for item in opf.iter():
        if item.tag.endswith("}item"):
            manifest[item.attrib["id"]] = item.attrib["href"]
    order = [manifest[i.attrib["idref"]] for i in opf.iter() if i.tag.endswith("}itemref") and i.attrib["idref"] in manifest]
    segs = []
    for n, href in enumerate(order, 1):
        try:
            html = z.read(base + href).decode("utf-8", errors="ignore")
        except KeyError:
            continue
        _, text = html_to_text(html)
        if len(text) > 40:
            parts = sectionize(text, 3000, prefix=f"Ch.{n}.")
            segs.extend(parts)
        progress(n, len(order))
    return segs


def extract_text(path, progress):
    text = read_text_file(path)
    if path.suffix.lower() == ".json":
        try:
            text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        except Exception:
            pass
    progress(1, 1)
    return sectionize(text)


def extract_image(path, progress):
    mime = "image/jpeg" if path.suffix.lower() in {".jpg", ".jpeg"} else "image/png"
    data = path.read_bytes()
    if path.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
        from PIL import Image

        buf = io.BytesIO()
        Image.open(path).convert("RGB").save(buf, "PNG")
        data, mime = buf.getvalue(), "image/png"
    text = llm.ocr_image(data, mime)
    progress(1, 1)
    return [("image", text)] if text else []


HANDLERS = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "pptx": extract_pptx,
    "xlsx": extract_xlsx,
    "csv": extract_csv,
    "html": extract_html,
    "epub": extract_epub,
    "text": extract_text,
    "code": extract_text,
    "image": extract_image,
}


def extract(path, progress=lambda a, b: None):
    kind = kind_of(path)
    if kind is None:
        raise ValueError("unsupported file type")
    return kind, HANDLERS[kind](path, progress)


def fetch_url(url, dest_dir):
    headers = {"User-Agent": "Mozilla/5.0 (Dars study assistant)"}
    with httpx.Client(follow_redirects=True, timeout=40, headers=headers) as c:
        r = c.get(url)
        r.raise_for_status()
    ctype = r.headers.get("content-type", "").lower()
    if "pdf" in ctype or url.lower().split("?")[0].endswith(".pdf"):
        name = (url.rstrip("/").split("/")[-1].split("?")[0] or "document") 
        name = name if name.lower().endswith(".pdf") else name + ".pdf"
        target = dest_dir / name
        target.write_bytes(r.content)
        return name, target
    title, text = html_to_text(r.text)
    name = (title or url.split("//")[-1].split("/")[0])[:90]
    safe = re.sub(r"[^\w؀-ۿ\- ]+", "", name).strip() or "page"
    target = dest_dir / f"{safe}.html"
    target.write_text(r.text, encoding="utf-8")
    return safe + ".html", target
